"""Multi-format document parsing.

Each parser turns a raw uploaded file into one or more LlamaIndex ``Document``
objects, attaching structural metadata (page / sheet / row / email field) that
later becomes the citation anchor and the Qdrant payload for scoped retrieval.

Supported (MVP): .pdf, .docx, .xlsx, .txt, .eml, images (.jpg/.jpeg/.png/.tiff/.bmp)
Explicitly rejected: .doc, .msg, .pst, .ost — no reliable pure-Python parser
without external tooling (LibreOffice headless / extract-msg).

Scanned PDF pages (no text layer) and image uploads are read via Tesseract OCR
(see services/ocr.py). For PDFs the ``ocr_mode`` argument controls this:
  * "auto"   — OCR a page only when its text layer is below a char threshold
  * "always" — OCR every page regardless of any text layer
  * "off"    — never OCR (text layer only; images then yield nothing)
"""
from __future__ import annotations

import email
from datetime import datetime, timezone
from email.message import Message
from pathlib import Path

import fitz  # pymupdf
import pandas as pd
from docx import Document as DocxDocument
from llama_index.core.schema import Document

from app.config import get_settings
from app.services import ocr

IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".tiff", ".bmp"}
SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".xlsx", ".txt", ".eml"} | IMAGE_EXTENSIONS
REJECTED_EXTENSIONS = {".doc", ".msg", ".pst", ".ost"}
OCR_MODES = {"auto", "always", "off"}


class UnsupportedFileTypeError(ValueError):
    """Raised for a known-but-unsupported format (clear message for the API)."""


def parse_file(
    file_path: str | Path,
    filename: str,
    claim_id: str,
    document_id: str,
    ocr_mode: str = "auto",
) -> list[Document]:
    """Dispatch on file extension and return LlamaIndex Documents.

    Every returned Document carries base metadata (source_filename, file_type,
    claim_id, document_id, upload_timestamp) plus a format-specific location key.

    ``ocr_mode`` ("auto" | "always" | "off") governs OCR of scanned PDF pages and
    is threaded through email attachments so nested PDFs/images OCR too.
    """
    ext = Path(filename).suffix.lower()

    if ext in REJECTED_EXTENSIONS:
        raise UnsupportedFileTypeError(
            f"'{ext}' is a legacy binary format not supported in this build. "
            f"Please convert to one of: {', '.join(sorted(SUPPORTED_EXTENSIONS))}."
        )
    if ext not in SUPPORTED_EXTENSIONS:
        raise UnsupportedFileTypeError(
            f"Unsupported file type '{ext}'. "
            f"Supported: {', '.join(sorted(SUPPORTED_EXTENSIONS))}."
        )

    base_meta = {
        "source_filename": filename,
        "file_type": ext.lstrip("."),
        "claim_id": claim_id,
        "document_id": document_id,
        "upload_timestamp": datetime.now(timezone.utc).isoformat(),
    }

    if ext == ".pdf":
        return _parse_pdf(file_path, base_meta, ocr_mode)
    if ext in IMAGE_EXTENSIONS:
        return _parse_image(file_path, base_meta, ocr_mode)
    if ext == ".docx":
        return _parse_docx(file_path, base_meta)
    if ext == ".xlsx":
        return _parse_xlsx(file_path, base_meta)
    if ext == ".eml":
        return _parse_eml(file_path, base_meta, ocr_mode)
    return _parse_txt(file_path, base_meta)


def _parse_pdf(file_path: str | Path, base_meta: dict, ocr_mode: str) -> list[Document]:
    """One Document per page so citations can reference a page number.

    Pages are processed in order (``enumerate``) and each carries its own
    ``page_label``, so OCR — which happens inline for a page that needs it —
    never changes page ordering. ``extraction_method`` records whether a page's
    text came from the native layer or from OCR.
    """
    settings = get_settings()
    docs: list[Document] = []
    with fitz.open(file_path) as pdf:
        for page_index, page in enumerate(pdf):
            text = page.get_text("text").strip()
            method = "text"

            if _should_ocr_page(text, ocr_mode, settings):
                ocr_text = ocr.ocr_page(page, dpi=settings.ocr_dpi, lang=settings.ocr_lang)
                # Prefer OCR only when it actually recovered more than the native
                # layer (e.g. an empty/sparse page); otherwise keep the real text.
                if len(ocr_text) > len(text):
                    text, method = ocr_text, "ocr"

            if not text:
                continue
            docs.append(
                Document(
                    text=text,
                    metadata={
                        **base_meta,
                        "page_label": page_index + 1,
                        "extraction_method": method,
                    },
                )
            )
    return docs


def _should_ocr_page(text: str, ocr_mode: str, settings) -> bool:
    """Decide whether a PDF page should be OCR'd, given its native text."""
    if ocr_mode == "off":
        return False
    if ocr_mode == "always":
        return True
    # "auto": OCR only pages whose text layer is empty or too sparse to trust.
    return settings.ocr_enabled and len(text) < settings.ocr_min_chars


def _parse_image(file_path: str | Path, base_meta: dict, ocr_mode: str) -> list[Document]:
    """Image uploads (e.g. a camera photo) have no text layer — OCR is the only
    way to read them, so we always OCR unless ``ocr_mode`` explicitly disables it.
    """
    if ocr_mode == "off":
        return []
    settings = get_settings()
    text = ocr.ocr_image_file(file_path, lang=settings.ocr_lang)
    if not text:
        return []
    return [
        Document(
            text=text,
            metadata={**base_meta, "page_label": 1, "extraction_method": "ocr"},
        )
    ]


def _parse_docx(file_path: str | Path, base_meta: dict) -> list[Document]:
    """Whole-document text (python-docx exposes no reliable page boundaries)."""
    docx = DocxDocument(str(file_path))
    paragraphs = [p.text for p in docx.paragraphs if p.text.strip()]
    text = "\n".join(paragraphs).strip()
    if not text:
        return []
    return [Document(text=text, metadata={**base_meta})]


def _parse_xlsx(file_path: str | Path, base_meta: dict) -> list[Document]:
    """One Document per non-empty row, preserving column context.

    Keeping each row as its own Document means chunking never truncates a policy
    schedule row mid-way, and citations can point at an exact sheet + row.
    """
    docs: list[Document] = []
    sheets = pd.read_excel(file_path, sheet_name=None, dtype=str)  # all sheets
    for sheet_name, df in sheets.items():
        df = df.fillna("")
        for row_pos, (_, row) in enumerate(df.iterrows()):
            cells = [f"{col}: {val}" for col, val in row.items() if str(val).strip()]
            if not cells:
                continue
            text = f"Sheet '{sheet_name}' | " + " | ".join(cells)
            docs.append(
                Document(
                    text=text,
                    metadata={
                        **base_meta,
                        "sheet_name": sheet_name,
                        # +2: 1 for header row, 1 for 1-based spreadsheet rows.
                        "row_number": row_pos + 2,
                    },
                )
            )
    return docs


def _parse_eml(file_path: str | Path, base_meta: dict, ocr_mode: str) -> list[Document]:
    """Email headers + body, recursively parsing supported attachments."""
    with open(file_path, "rb") as fh:
        msg: Message = email.message_from_binary_file(fh)

    headers = {
        "From": msg.get("From", ""),
        "To": msg.get("To", ""),
        "Subject": msg.get("Subject", ""),
        "Date": msg.get("Date", ""),
    }
    header_text = "\n".join(f"{k}: {v}" for k, v in headers.items() if v)
    body_text = _extract_email_body(msg)

    docs: list[Document] = []
    combined = "\n\n".join(part for part in [header_text, body_text] if part).strip()
    if combined:
        docs.append(
            Document(
                text=combined,
                metadata={**base_meta, "email_field": "body"},
            )
        )

    docs.extend(_parse_email_attachments(msg, base_meta, ocr_mode))
    return docs


def _extract_email_body(msg: Message) -> str:
    if msg.is_multipart():
        parts = []
        for part in msg.walk():
            if part.get_content_type() == "text/plain" and not part.get_filename():
                payload = part.get_payload(decode=True)
                if payload:
                    parts.append(payload.decode(part.get_content_charset() or "utf-8", "replace"))
        return "\n".join(parts).strip()
    payload = msg.get_payload(decode=True)
    if payload:
        return payload.decode(msg.get_content_charset() or "utf-8", "replace").strip()
    return ""


def _parse_email_attachments(
    msg: Message, base_meta: dict, ocr_mode: str
) -> list[Document]:
    """Recurse supported attachments through the main parser via temp files."""
    import tempfile

    docs: list[Document] = []
    for part in msg.walk():
        filename = part.get_filename()
        if not filename:
            continue
        ext = Path(filename).suffix.lower()
        if ext not in SUPPORTED_EXTENSIONS:
            continue
        payload = part.get_payload(decode=True)
        if not payload:
            continue
        with tempfile.NamedTemporaryFile(suffix=ext, delete=True) as tmp:
            tmp.write(payload)
            tmp.flush()
            attached = parse_file(
                tmp.name,
                filename,
                base_meta["claim_id"],
                base_meta["document_id"],
                ocr_mode,
            )
        for d in attached:
            d.metadata["attached_to"] = base_meta["source_filename"]
        docs.extend(attached)
    return docs


def _parse_txt(file_path: str | Path, base_meta: dict) -> list[Document]:
    text = Path(file_path).read_text(encoding="utf-8", errors="replace").strip()
    if not text:
        return []
    return [Document(text=text, metadata={**base_meta})]
